# -*- coding: utf-8 -*-
"""
地基雷达监测报告生成器（双数据源GUI版）- V4 修正版
- 文件路径显示完整，易于区分
- 短时数据用于6小时报告，长时数据用于24小时和周报
- 自动推荐时间节点，导出TXT和Excel
- V4：24h正文不输出周对比；修正Excel总累积位移；增强6h时间选择和字段校验
"""
import os
# 如果遇到 Tcl/Tk 错误，请取消注释下面两行并修改为您的实际路径
# os.environ['TCL_LIBRARY'] = r'C:\path\to\tcl8.6'
# os.environ['TK_LIBRARY'] = r'C:\path\to\tk8.6'

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
from datetime import datetime, timedelta
from collections import OrderedDict

# ---------- 数据加载（不变） ----------
def _clean_header(header):
    if header is None:
        return None
    return str(header).strip()

def load_data(filepath, is_point=False):
    wb = openpyxl.load_workbook(filepath)
    ws = wb['位移']
    headers = [_clean_header(cell.value) for cell in ws[1]]
    col_map = {}
    for i, h in enumerate(headers[1:], 1):
        if h:
            col_map[h] = i
    data = OrderedDict()
    start_row = 5 if is_point else 2
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True):
        time_raw = row[0]
        if time_raw is None or str(time_raw).strip() == '':
            continue
        ts = str(time_raw).strip()
        if is_point and ts in ('X', 'Y', 'Z'):
            continue
        dt = _parse_time(ts)
        if dt is None:
            continue
        values = {}
        for name, col_idx in col_map.items():
            v = row[col_idx]
            values[name] = v if v is not None else 0.0
        data[ts] = values
    wb.close()
    return data, col_map

def _parse_time(s):
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H'):
        try:
            return datetime.strptime(s, fmt)
        except:
            pass
    return None

def parse_time(s):
    dt = _parse_time(s)
    if dt is None:
        raise ValueError(f"无法解析时间: {s}")
    return dt

def fmt_val(val):
    return f"+{val:.1f}" if val >= 0 else f"{val:.1f}"

def compute_item_disp(data, s, e, name):
    sr = data.get(s)
    er = data.get(e)
    if sr is None or er is None:
        return None, None
    return (er.get(name, 0.0) or 0.0) - (sr.get(name, 0.0) or 0.0), (er.get(name, 0.0) or 0.0)

def compute_group_disp(data, s, e, names):
    sr = data.get(s)
    er = data.get(e)
    if sr is None or er is None:
        return None, None, None, None
    d = {}
    for n in names:
        sv = sr.get(n, 0.0) or 0.0
        ev = er.get(n, 0.0) or 0.0
        d[n] = (ev - sv, ev)
    if not d:
        return None, None, None, None
    mx = max(d, key=lambda k: d[k][0])
    mn = min(d, key=lambda k: d[k][0])
    return d[mn][0], d[mx][0], mx, d[mx][1]

def get_sorted_names(d):
    names = list(d.keys())
    def sk(n):
        np = ''.join(c for c in n if c.isdigit())
        return int(np) if np else 999
    return sorted(names, key=sk)

# ---------- 分组定义 ----------
FACE_GROUPS = OrderedDict([
    ('上部及后缘(全区变形峰值区)', {
        'face_names': ['M6', 'M7', 'M8'],
        'face_range': 'M6-M8',
        'point_names': ['S6', 'S7', 'S8'],
        'point_range': 'S6-S8',
    }),
    ('中部公路边坡', {
        'face_names': ['M1', 'M2', 'M3', 'M4', 'M5'],
        'face_range': 'M1-M5',
        'point_names': ['S1', 'S2', 'S3', 'S4', 'S5'],
        'point_range': 'S1-S5',
    }),
    ('下部边坡', {
        'face_names': ['M9', 'M10', 'M11', 'M12', 'M13', 'M14'],
        'face_range': 'M9-M14',
        'point_names': ['S9', 'S10', 'S11', 'S12', 'S13', 'S14'],
        'point_range': 'S9-S14',
    }),
])

REQUIRED_FACE_NAMES = sorted({n for g in FACE_GROUPS.values() for n in g['face_names']}, key=lambda x: int(''.join(c for c in x if c.isdigit())))
REQUIRED_POINT_NAMES = sorted({n for g in FACE_GROUPS.values() for n in g['point_names']}, key=lambda x: int(''.join(c for c in x if c.isdigit())))

def validate_required_columns(col_map, required, label):
    missing = [n for n in required if n not in col_map]
    if missing:
        raise ValueError(f"{label}缺少必要列：{'、'.join(missing)}")

# ---------- GUI ----------
class App:
    def __init__(self, root):
        self.root = root
        root.title("地基雷达监测报告生成器（双数据源）")
        root.geometry("1100x980")
        root.minsize(950, 800)

        # 数据存储
        self.face_short = None
        self.point_short = None
        self.face_long = None
        self.point_long = None
        self.face_short_cols = None
        self.point_short_cols = None
        self.face_long_cols = None
        self.point_long_cols = None
        self.short_times = []
        self.long_times = []
        self.face_short_file = ''
        self.point_short_file = ''
        self.face_long_file = ''
        self.point_long_file = ''

        style = ttk.Style()
        style.theme_use('clam')

        # 标题
        header = ttk.Frame(root)
        header.pack(fill='x', padx=12, pady=(12, 0))
        ttk.Label(header, text="地基雷达监测报告生成器（双数据源）",
                  font=('Microsoft YaHei', 16, 'bold')).pack()
        ttk.Label(header, text="短时数据（小时级）用于6小时报告；长时数据（日级）用于24小时和周报",
                  font=('Microsoft YaHei', 10), foreground='gray').pack(pady=(2, 0))

        # ---- 文件选择（优化布局） ----
        f_frame = ttk.LabelFrame(root, text=" 数据文件 ", padding=12)
        f_frame.pack(fill='x', padx=12, pady=10)

        # 短时区域
        ttk.Label(f_frame, text="【短时数据（小时级）】", font=('Microsoft YaHei', 10, 'bold')).grid(
            row=0, column=0, columnspan=3, sticky='w', pady=(0,5))

        # 短时面
        ttk.Label(f_frame, text="短时面:", width=12).grid(row=1, column=0, sticky='w', padx=2)
        self.face_short_path = tk.StringVar()
        entry_face_short = ttk.Entry(f_frame, textvariable=self.face_short_path, state='readonly')
        entry_face_short.grid(row=1, column=1, sticky='ew', padx=5)
        ttk.Button(f_frame, text="浏览...", command=lambda: self.browse('face_short'), width=8).grid(row=1, column=2, padx=2)

        # 短时点
        ttk.Label(f_frame, text="短时点:", width=12).grid(row=2, column=0, sticky='w', padx=2)
        self.point_short_path = tk.StringVar()
        entry_point_short = ttk.Entry(f_frame, textvariable=self.point_short_path, state='readonly')
        entry_point_short.grid(row=2, column=1, sticky='ew', padx=5)
        ttk.Button(f_frame, text="浏览...", command=lambda: self.browse('point_short'), width=8).grid(row=2, column=2, padx=2)

        # 分隔线
        ttk.Separator(f_frame, orient='horizontal').grid(row=3, column=0, columnspan=3, sticky='ew', pady=8)

        # 长时区域
        ttk.Label(f_frame, text="【长时数据（日级）】", font=('Microsoft YaHei', 10, 'bold')).grid(
            row=4, column=0, columnspan=3, sticky='w', pady=(0,5))

        # 长时面
        ttk.Label(f_frame, text="长时面:", width=12).grid(row=5, column=0, sticky='w', padx=2)
        self.face_long_path = tk.StringVar()
        entry_face_long = ttk.Entry(f_frame, textvariable=self.face_long_path, state='readonly')
        entry_face_long.grid(row=5, column=1, sticky='ew', padx=5)
        ttk.Button(f_frame, text="浏览...", command=lambda: self.browse('face_long'), width=8).grid(row=5, column=2, padx=2)

        # 长时点
        ttk.Label(f_frame, text="长时点:", width=12).grid(row=6, column=0, sticky='w', padx=2)
        self.point_long_path = tk.StringVar()
        entry_point_long = ttk.Entry(f_frame, textvariable=self.point_long_path, state='readonly')
        entry_point_long.grid(row=6, column=1, sticky='ew', padx=5)
        ttk.Button(f_frame, text="浏览...", command=lambda: self.browse('point_long'), width=8).grid(row=6, column=2, padx=2)

        # 设置列权重，使Entry列伸展
        f_frame.columnconfigure(1, weight=1)

        # 加载按钮
        bf = ttk.Frame(f_frame)
        bf.grid(row=7, column=0, columnspan=3, pady=8)
        ttk.Button(bf, text="加载全部数据", command=self.load_all_data, width=15).pack()
        self.status_label = ttk.Label(bf, text="", foreground='gray')
        self.status_label.pack(pady=4)

        # ---- Notebook ----
        self.notebook = ttk.Notebook(root)
        self.notebook.pack(fill='both', expand=True, padx=12, pady=(0, 10))

        self.tab6 = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.tab6, text="  6小时  ")
        self.build_tab6()

        self.tab24 = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.tab24, text="  24小时  ")
        self.build_tab24()

        self.tab7 = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.tab7, text="  周报  ")
        self.build_tab7()

    # ----- 文件浏览 -----
    def browse(self, tp):
        fp = filedialog.askopenfilename(filetypes=[("Excel files","*.xlsx")])
        if fp:
            if tp == 'face_short':
                self.face_short_path.set(fp); self.face_short_file = fp
            elif tp == 'point_short':
                self.point_short_path.set(fp); self.point_short_file = fp
            elif tp == 'face_long':
                self.face_long_path.set(fp); self.face_long_file = fp
            elif tp == 'point_long':
                self.point_long_path.set(fp); self.point_long_file = fp

    # ----- 加载数据 -----
    def load_all_data(self):
        if not (self.face_short_file and self.point_short_file and self.face_long_file and self.point_long_file):
            messagebox.showwarning("提示", "请选择所有四个数据文件")
            return
        try:
            fd_s, fc_s = load_data(self.face_short_file, is_point=False)
            pd_s, pc_s = load_data(self.point_short_file, is_point=True)
            validate_required_columns(fc_s, REQUIRED_FACE_NAMES, "短时面数据")
            validate_required_columns(pc_s, REQUIRED_POINT_NAMES, "短时点数据")
            self.face_short = OrderedDict(sorted(fd_s.items(), key=lambda x: parse_time(x[0])))
            self.point_short = OrderedDict(sorted(pd_s.items(), key=lambda x: parse_time(x[0])))
            self.face_short_cols = fc_s; self.point_short_cols = pc_s
            self.short_times = sorted(set(self.face_short.keys()) & set(self.point_short.keys()), key=lambda x: parse_time(x))

            fd_l, fc_l = load_data(self.face_long_file, is_point=False)
            pd_l, pc_l = load_data(self.point_long_file, is_point=True)
            validate_required_columns(fc_l, REQUIRED_FACE_NAMES, "长时面数据")
            validate_required_columns(pc_l, REQUIRED_POINT_NAMES, "长时点数据")
            self.face_long = OrderedDict(sorted(fd_l.items(), key=lambda x: parse_time(x[0])))
            self.point_long = OrderedDict(sorted(pd_l.items(), key=lambda x: parse_time(x[0])))
            self.face_long_cols = fc_l; self.point_long_cols = pc_l
            self.long_times = sorted(set(self.face_long.keys()) & set(self.point_long.keys()), key=lambda x: parse_time(x))

            if not self.short_times or not self.long_times:
                messagebox.showerror("错误", "短时或长时数据没有共同时间点")
                return

            self.refresh_selects()
            self.auto_6h()
            self.auto_24h()
            self.auto_7d()
            self.status_label.config(
                text=f"短时: {len(self.short_times)} 点，长时: {len(self.long_times)} 点",
                foreground='green')
        except Exception as e:
            messagebox.showerror("错误", f"加载失败: {e}")
            import traceback; traceback.print_exc()

    # ----- 构建标签页（不变） -----
    def build_tab6(self):
        tf = ttk.Frame(self.tab6); tf.pack(fill='x')
        ttk.Label(tf, text="开始时间:").pack(side='left')
        self.s6_cb = ttk.Combobox(tf, state='readonly', width=28); self.s6_cb.pack(side='left', padx=5)
        ttk.Label(tf, text="结束时间:").pack(side='left', padx=(15,0))
        self.e6_cb = ttk.Combobox(tf, state='readonly', width=28); self.e6_cb.pack(side='left', padx=5)
        ttk.Button(tf, text="自动推荐", command=self.auto_6h).pack(side='left', padx=10)
        ttk.Button(tf, text="生成报告", command=self.gen_6h).pack(side='left')
        self.r6_text = tk.Text(self.tab6, height=12, wrap='word', font=('Microsoft YaHei', 10), state='disabled')
        self.r6_text.pack(fill='both', expand=True, pady=10)
        bf6 = ttk.Frame(self.tab6); bf6.pack(fill='x')
        ttk.Button(bf6, text="下载 .txt", command=lambda: self.dl_txt('6h')).pack(side='left', padx=5)

    def build_tab24(self):
        cf = ttk.LabelFrame(self.tab24, text=" 当前周期 ", padding=8); cf.pack(fill='x')
        r1 = ttk.Frame(cf); r1.pack(fill='x', pady=3)
        ttk.Label(r1, text="开始时间:").pack(side='left')
        self.s24_cb = ttk.Combobox(r1, state='readonly', width=26); self.s24_cb.pack(side='left', padx=5)
        ttk.Label(r1, text="结束时间:").pack(side='left', padx=(10,0))
        self.e24_cb = ttk.Combobox(r1, state='readonly', width=26); self.e24_cb.pack(side='left', padx=5)
        pf = ttk.LabelFrame(self.tab24, text=" Excel对比周期（仅用于对比表，不写入日报正文） ", padding=8); pf.pack(fill='x', pady=8)
        r2 = ttk.Frame(pf); r2.pack(fill='x', pady=3)
        ttk.Label(r2, text="开始时间:").pack(side='left')
        self.ps24_cb = ttk.Combobox(r2, state='readonly', width=26); self.ps24_cb.pack(side='left', padx=5)
        ttk.Label(r2, text="结束时间:").pack(side='left', padx=(10,0))
        self.pe24_cb = ttk.Combobox(r2, state='readonly', width=26); self.pe24_cb.pack(side='left', padx=5)
        bf24 = ttk.Frame(self.tab24); bf24.pack(fill='x', pady=5)
        ttk.Button(bf24, text="自动推荐", command=self.auto_24h).pack(side='left', padx=5)
        ttk.Button(bf24, text="生成报告", command=self.gen_24h).pack(side='left', padx=5)
        self.r24_text = tk.Text(self.tab24, height=12, wrap='word', font=('Microsoft YaHei', 10), state='disabled')
        self.r24_text.pack(fill='both', expand=True, pady=8)
        bf24b = ttk.Frame(self.tab24); bf24b.pack(fill='x')
        ttk.Button(bf24b, text="下载 .txt", command=lambda: self.dl_txt('24h')).pack(side='left', padx=5)
        ttk.Button(bf24b, text="下载 Excel对比表", command=self.dl_excel_24h).pack(side='left', padx=5)

    def build_tab7(self):
        cf = ttk.LabelFrame(self.tab7, text=" 当前周期（7天） ", padding=8); cf.pack(fill='x')
        r1 = ttk.Frame(cf); r1.pack(fill='x', pady=3)
        ttk.Label(r1, text="开始时间:").pack(side='left')
        self.s7_cb = ttk.Combobox(r1, state='readonly', width=26); self.s7_cb.pack(side='left', padx=5)
        ttk.Label(r1, text="结束时间:").pack(side='left', padx=(10,0))
        self.e7_cb = ttk.Combobox(r1, state='readonly', width=26); self.e7_cb.pack(side='left', padx=5)
        pf = ttk.LabelFrame(self.tab7, text=" 对比周期（前一周） ", padding=8); pf.pack(fill='x', pady=8)
        r2 = ttk.Frame(pf); r2.pack(fill='x', pady=3)
        ttk.Label(r2, text="开始时间:").pack(side='left')
        self.ps7_cb = ttk.Combobox(r2, state='readonly', width=26); self.ps7_cb.pack(side='left', padx=5)
        ttk.Label(r2, text="结束时间:").pack(side='left', padx=(10,0))
        self.pe7_cb = ttk.Combobox(r2, state='readonly', width=26); self.pe7_cb.pack(side='left', padx=5)
        bf7 = ttk.Frame(self.tab7); bf7.pack(fill='x', pady=5)
        ttk.Button(bf7, text="自动推荐（周五节点）", command=self.auto_7d).pack(side='left', padx=5)
        ttk.Button(bf7, text="生成报告", command=self.gen_7d).pack(side='left', padx=5)
        self.r7_text = tk.Text(self.tab7, height=12, wrap='word', font=('Microsoft YaHei', 10), state='disabled')
        self.r7_text.pack(fill='both', expand=True, pady=8)
        bf7b = ttk.Frame(self.tab7); bf7b.pack(fill='x')
        ttk.Button(bf7b, text="下载 .txt", command=lambda: self.dl_txt('7d')).pack(side='left', padx=5)
        ttk.Button(bf7b, text="下载 Excel对比表", command=self.dl_excel_7d).pack(side='left', padx=5)

    # ----- 辅助方法 -----
    def refresh_selects(self):
        for cb in [self.s6_cb, self.e6_cb]:
            cb['values'] = self.short_times
        for cb in [self.s24_cb, self.e24_cb, self.ps24_cb, self.pe24_cb,
                   self.s7_cb, self.e7_cb, self.ps7_cb, self.pe7_cb]:
            cb['values'] = self.long_times

    def set_cb(self, cb, val):
        if val and val in cb['values']:
            cb.set(val)

    def _find_nearest(self, times, target_hour, target_min=0, before=True, target_date=None):
        best = None
        for t in times:
            dt = parse_time(t)
            if target_date is not None and dt.date() != target_date:
                continue
            if before and (dt.hour > target_hour or (dt.hour == target_hour and dt.minute > target_min)):
                continue
            if not before and (dt.hour < target_hour or (dt.hour == target_hour and dt.minute < target_min)):
                continue
            best = t
        return best

    def _find_nearest_time(self, times, ref_time):
        if not times:
            return None
        ref_dt = parse_time(ref_time)
        best = min(times, key=lambda t: abs((parse_time(t) - ref_dt).total_seconds()))
        return best

    # ----- 自动推荐 -----
    def auto_6h(self):
        if not self.short_times:
            return
        latest = parse_time(self.short_times[-1])
        sh = (latest.hour // 6) * 6
        current_start = latest.replace(hour=sh, minute=0, second=0, microsecond=0)

        # V4：按真实 6 小时时间桶向前寻找，避免只比较小时数造成跨日期误选。
        for back in range(8):
            start_dt = current_start - timedelta(hours=6 * back)
            end_dt = start_dt + timedelta(hours=6)
            bucket = [t for t in self.short_times
                      if start_dt <= parse_time(t) < end_dt
                      and (back > 0 or parse_time(t) <= latest)]
            if len(bucket) >= 2:
                self.set_cb(self.s6_cb, bucket[0])
                self.set_cb(self.e6_cb, bucket[-1])
                return

        # 极端缺数时保留全部可用范围供手动选择。
        self.set_cb(self.s6_cb, self.short_times[0])
        self.set_cb(self.e6_cb, self.short_times[-1])

    def auto_24h(self):
        if not self.long_times: return
        today = parse_time(self.long_times[-1]).date()
        end = self._find_nearest(self.long_times, 15, before=True, target_date=today)
        if end is None: end = self.long_times[-1]
        yesterday = today - timedelta(days=1)
        start = self._find_nearest(self.long_times, 15, before=True, target_date=yesterday)
        if start is None: start = self.long_times[0]
        self.set_cb(self.s24_cb, start); self.set_cb(self.e24_cb, end)
        pe = start
        prev_day = yesterday - timedelta(days=1)
        ps = self._find_nearest(self.long_times, 15, before=True, target_date=prev_day)
        if ps is None: ps = self.long_times[0]
        self.set_cb(self.ps24_cb, ps); self.set_cb(self.pe24_cb, pe)

    def auto_7d(self):
        if not self.long_times: return
        latest = parse_time(self.long_times[-1])
        found = None
        for i in range(14):
            day = latest - timedelta(days=i)
            if day.weekday() == 4:
                found = day; break
        if found is None: return
        end = self._find_nearest(self.long_times, 15, before=True, target_date=found)
        if end is None: end = self.long_times[-1]
        start_date = found - timedelta(days=7)
        start = self._find_nearest(self.long_times, 15, before=True, target_date=start_date)
        if start is None: start = self.long_times[0]
        ps_date = start_date - timedelta(days=7)
        ps = self._find_nearest(self.long_times, 15, before=True, target_date=ps_date)
        if ps is None: ps = self.long_times[0]
        pe = start
        self.set_cb(self.s7_cb, start); self.set_cb(self.e7_cb, end)
        self.set_cb(self.ps7_cb, ps); self.set_cb(self.pe7_cb, pe)

    # ----- 报告生成核心（修正总累积位移获取） -----
    def gen_report_text(self, ss, es, label, disp_face, disp_point, total_face, total_point,
                        compare=None, display_start=None, display_end=None, compare_label="与上周相比"):
        if display_start is None: display_start = ss
        if display_end is None: display_end = es
        lines = [f"地基雷达{label}监测结果（{display_start} 至 {display_end}）"]
        for gn, gi in FACE_GROUPS.items():
            f_min, f_max, fn, ft_disp = compute_group_disp(disp_face, ss, es, gi['face_names'])
            p_min, p_max, pn, pt_disp = compute_group_disp(disp_point, ss, es, gi['point_names'])
            if f_min is None or p_min is None: continue

            def get_total_value(data_dict, time_key, name):
                if time_key in data_dict:
                    return data_dict[time_key].get(name, 0.0) or 0.0
                nearest = self._find_nearest_time(list(data_dict.keys()), time_key)
                if nearest is None:
                    return 0.0
                return data_dict[nearest].get(name, 0.0) or 0.0

            ft_total = get_total_value(total_face, es, fn)
            pt_total = get_total_value(total_point, es, pn)

            line = (f"- {gn}："
                    f"监测面（{gi['face_range']}）"
                    f"{label}时段位移{fmt_val(f_min)}~{fmt_val(f_max)}mm，"
                    f"最大值位于{fn}，截止期末总累积位移{fmt_val(ft_total)}mm；"
                    f"监测点（{gi['point_range']}）"
                    f"{label}时段位移{fmt_val(p_min)}~{fmt_val(p_max)}mm，"
                    f"最大值位于{pn}，截止期末总累积位移{fmt_val(pt_total)}mm。")
            if compare:
                ps, pe = compare
                f_delta = {}
                for n in gi['face_names']:
                    d_this, _ = compute_item_disp(disp_face, ss, es, n)
                    d_last, _ = compute_item_disp(disp_face, ps, pe, n)
                    if d_this is not None and d_last is not None:
                        f_delta[n] = d_this - d_last
                if f_delta:
                    f_delta_min = min(f_delta.values()); f_delta_max = max(f_delta.values())
                    p_delta = {}
                    for n in gi['point_names']:
                        d_this, _ = compute_item_disp(disp_point, ss, es, n)
                        d_last, _ = compute_item_disp(disp_point, ps, pe, n)
                        if d_this is not None and d_last is not None:
                            p_delta[n] = d_this - d_last
                    if p_delta:
                        p_delta_min = min(p_delta.values()); p_delta_max = max(p_delta.values())
                        line += (f" {compare_label}，监测面位移变化{fmt_val(f_delta_min)}~{fmt_val(f_delta_max)}mm，"
                                 f"监测点位移变化{fmt_val(p_delta_min)}~{fmt_val(p_delta_max)}mm。")
                    else:
                        line += f" {compare_label}，监测面位移变化{fmt_val(f_delta_min)}~{fmt_val(f_delta_max)}mm。"
            lines.append(line)
        return "\n".join(lines) if len(lines) > 1 else None

    def _round_to_hour(self, t_str):
        dt = parse_time(t_str)
        if dt.minute >= 30: dt += timedelta(hours=1)
        return dt.replace(minute=0, second=0).strftime('%Y-%m-%d %H:%M:%S')

    def _to_15_hour(self, t_str):
        dt = parse_time(t_str)
        return dt.replace(hour=15, minute=0, second=0).strftime('%Y-%m-%d %H:%M:%S')

    # ----- 生成报告 -----
    def gen_6h(self):
        if self.face_short is None:
            messagebox.showwarning("提示", "请先加载短时数据"); return
        ss = self.s6_cb.get(); es = self.e6_cb.get()
        if not ss or not es:
            messagebox.showwarning("提示", "请选择起止时间"); return
        if parse_time(ss) >= parse_time(es):
            messagebox.showwarning("提示", "开始时间必须早于结束时间"); return
        disp_start = self._round_to_hour(ss); disp_end = self._round_to_hour(es)
        text = self.gen_report_text(ss, es, '6小时',
                                    disp_face=self.face_short, disp_point=self.point_short,
                                    total_face=self.face_long, total_point=self.point_long,
                                    display_start=disp_start, display_end=disp_end)
        if text:
            self._show_text(self.r6_text, text); self._last6 = ('6h', text, ss, es)

    def gen_24h(self):
        if self.face_long is None:
            messagebox.showwarning("提示", "请先加载长时数据"); return
        ss = self.s24_cb.get(); es = self.e24_cb.get()
        ps = self.ps24_cb.get(); pe = self.pe24_cb.get()
        if not ss or not es:
            messagebox.showwarning("提示", "请选择当前周期"); return
        if parse_time(ss) >= parse_time(es):
            messagebox.showwarning("提示", "开始时间必须早于结束时间"); return
        has_prev = bool(ps and pe and parse_time(ps) < parse_time(pe))
        disp_start = self._to_15_hour(ss); disp_end = self._to_15_hour(es)
        # V4：24小时日报正文仅报告当前周期，不输出任何“与上周相比”或前日对比文字。
        # ps/pe 仍保留给 Excel 对比表使用。
        text = self.gen_report_text(ss, es, '24小时',
                                    disp_face=self.face_long, disp_point=self.point_long,
                                    total_face=self.face_long, total_point=self.point_long,
                                    compare=None,
                                    display_start=disp_start, display_end=disp_end)
        if text:
            self._show_text(self.r24_text, text)
            self._last24 = ('24h', text, ss, es, ps, pe, has_prev)

    def gen_7d(self):
        if self.face_long is None:
            messagebox.showwarning("提示", "请先加载长时数据"); return
        ss = self.s7_cb.get(); es = self.e7_cb.get()
        ps = self.ps7_cb.get(); pe = self.pe7_cb.get()
        if not ss or not es:
            messagebox.showwarning("提示", "请选择当前周期"); return
        if parse_time(ss) >= parse_time(es):
            messagebox.showwarning("提示", "开始时间必须早于结束时间"); return
        has_prev = bool(ps and pe and parse_time(ps) < parse_time(pe))
        disp_start = self._to_15_hour(ss); disp_end = self._to_15_hour(es)
        text = self.gen_report_text(ss, es, '7天',
                                    disp_face=self.face_long, disp_point=self.point_long,
                                    total_face=self.face_long, total_point=self.point_long,
                                    compare=(ps, pe) if has_prev else None,
                                    display_start=disp_start, display_end=disp_end,
                                    compare_label="与上周相比")
        if text:
            self._show_text(self.r7_text, text)
            self._last7 = ('7d', text, ss, es, ps, pe, has_prev)

    def _show_text(self, widget, text):
        widget.config(state='normal'); widget.delete('1.0', 'end')
        widget.insert('1.0', text); widget.config(state='disabled')

    # ----- Excel 保存 -----
    def _save_excel(self, fp, ss, es, ps, pe, has_prev, face_data, point_data, days=1):
        sd = parse_time(ss); ed = parse_time(es)
        dl = f"{sd.month}/{sd.day}-{ed.month}/{ed.day}"
        valid_prev = bool(has_prev and ps and pe and parse_time(ps) < parse_time(pe))
        headers = ['监测点/面', f'{dl}位移量(mm)']
        if valid_prev:
            psd = parse_time(ps); ped = parse_time(pe)
            pdl = f"{psd.month}/{psd.day}-{ped.month}/{ped.day}"
            headers.append(f'相对于{pdl}变化量(mm)')
        headers.append('总累积位移量(mm)')

        def total_at_end(data_dict, end_key, name):
            nearest = self._find_nearest_time(list(data_dict.keys()), end_key)
            if nearest is None:
                return 0.0
            # V4：源表本身就是“累计位移”，总累积位移应直接取期末值，不能再减首期值。
            return data_dict[nearest].get(name, 0.0) or 0.0

        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "面"; ws.append(headers)
        for n in get_sorted_names(self.face_long_cols):
            if n == 'W1':
                continue
            da, _ = compute_item_disp(face_data, ss, es, n)
            if da is None:
                continue
            row = [n, round(da, 1)]
            if valid_prev:
                db, _ = compute_item_disp(face_data, ps, pe, n)
                row.append(round(da - db, 1) if db is not None else '')
            row.append(round(total_at_end(face_data, es, n), 1))
            ws.append(row)

        ws2 = wb.create_sheet("点"); ws2.append(headers)
        for n in get_sorted_names(self.point_long_cols):
            da, _ = compute_item_disp(point_data, ss, es, n)
            if da is None:
                continue
            row = [n, round(da, 1)]
            if valid_prev:
                db, _ = compute_item_disp(point_data, ps, pe, n)
                row.append(round(da - db, 1) if db is not None else '')
            row.append(round(total_at_end(point_data, es, n), 1))
            ws2.append(row)
        wb.save(fp)

    # ----- 下载 -----
    def dl_txt(self, tp):
        if tp == '6h' and hasattr(self, '_last6'):
            label, text = self._last6[0], self._last6[1]
        elif tp == '24h' and hasattr(self, '_last24'):
            label, text = self._last24[0], self._last24[1]
        elif tp == '7d' and hasattr(self, '_last7'):
            label, text = self._last7[0], self._last7[1]
        else:
            messagebox.showwarning("提示", "请先生成对应报告"); return
        fp = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Text","*.txt")],
                                         initialfile=f"{tp}_report.txt")
        if fp:
            with open(fp, 'w', encoding='utf-8') as f:
                f.write(text + "\n")
            messagebox.showinfo("完成", f"已保存: {os.path.basename(fp)}")

    def dl_excel_24h(self):
        if not hasattr(self, '_last24') or self._last24[0] != '24h':
            messagebox.showwarning("提示", "请先生成24小时报告"); return
        _, _, ss, es, ps, pe, has_prev = self._last24
        fp = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")],
                                         initialfile="24h对比表.xlsx")
        if not fp: return
        self._save_excel(fp, ss, es, ps, pe, has_prev, self.face_long, self.point_long, days=1)
        messagebox.showinfo("完成", f"已保存: {os.path.basename(fp)}")

    def dl_excel_7d(self):
        if not hasattr(self, '_last7') or self._last7[0] != '7d':
            messagebox.showwarning("提示", "请先生成周报"); return
        _, _, ss, es, ps, pe, has_prev = self._last7
        fp = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel","*.xlsx")],
                                         initialfile="周报对比表.xlsx")
        if not fp: return
        self._save_excel(fp, ss, es, ps, pe, has_prev, self.face_long, self.point_long, days=7)
        messagebox.showinfo("完成", f"已保存: {os.path.basename(fp)}")

if __name__ == '__main__':
    root = tk.Tk()
    app = App(root)
    root.mainloop()